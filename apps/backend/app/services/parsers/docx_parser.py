import io
from typing import Dict, Any, List
import docx
from app.services.parsers.base import BaseParser, ParsedDocument, PageContent


class DocxParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)
        paragraphs: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        metadata: Dict[str, Any] = {
            "source": filename,
            "paragraph_count": len(paragraphs),
            "char_count": len(full_text),
        }
        pages = [PageContent(page_number=1, text=full_text)]
        return ParsedDocument(text=full_text, metadata=metadata, pages=pages)
